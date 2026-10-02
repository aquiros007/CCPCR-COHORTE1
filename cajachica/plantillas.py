"""Plantillas para los custodios y datos de demostración."""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .util import solo_digitos

AZUL = "1F3A5F"
COLUMNAS_FACTURAS = ["Fecha factura", "Tipo documento", "Clave numérica", "Número de factura", "Proveedor",
                     "Cédula proveedor", "Receptor (a nombre de)", "Cédula receptor", "Descripción del gasto",
                     "Categoría", "Subtotal", "IVA", "Total", "Moneda", "Solicitado por", "Autorizado por"]
ANCHOS_FACTURAS = [13, 11, 56, 24, 30, 15, 32, 15, 40, 22, 12, 11, 12, 8, 22, 22]


def _encabezado(ws, fila, columnas, anchos):
    for c, (h, w) in enumerate(zip(columnas, anchos), 1):
        celda = ws.cell(row=fila, column=c, value=h)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[celda.column_letter].width = w


def _hoja_encabezado(wb, campos):
    ws = wb.active
    ws.title = "Encabezado"
    _encabezado(ws, 1, ["Campo", "Valor"], [34, 40])
    for i, (campo, valor) in enumerate(campos, 2):
        ws.cell(row=i, column=1, value=campo).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valor)
    return ws


def libro_liquidacion(encabezado: list[tuple], facturas: list[list], politica: dict) -> Workbook:
    wb = Workbook()
    _hoja_encabezado(wb, encabezado)
    ws = wb.create_sheet("Facturas")
    _encabezado(ws, 1, COLUMNAS_FACTURAS, ANCHOS_FACTURAS)
    ws.freeze_panes = "A2"
    for r in range(2, 202):  # formato texto: evita que Excel dañe la clave de 50 dígitos
        for c in (3, 4, 6, 8):
            ws.cell(row=r, column=c).number_format = "@"
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        for c in (11, 12, 13):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    for r, fila in enumerate(facturas, 2):
        for c, v in enumerate(fila, 1):
            ws.cell(row=r, column=c, value=v)
    listas = wb.create_sheet("Listas")
    for i, cat in enumerate(politica.get("categorias_permitidas", []), 1):
        listas.cell(row=i, column=1, value=cat)
    for i, (cod, nom) in enumerate([("01", "Factura electrónica"), ("04", "Tiquete electrónico")], 1):
        listas.cell(row=i, column=2, value=cod)
        listas.cell(row=i, column=3, value=nom)
    for i, codigo in enumerate(sorted(politica.get("_catalogo") or {}), 1):
        listas.cell(row=i, column=4, value=codigo)
    listas.sheet_state = "hidden"
    n = max(1, len(politica.get("categorias_permitidas", [])))
    dv_cat = DataValidation(type="list", formula1=f"=Listas!$A$1:$A${n}", allow_blank=True)
    dv_tipo = DataValidation(type="list", formula1='"01,04"', allow_blank=True)
    dv_mon = DataValidation(type="list", formula1='"CRC,USD"', allow_blank=True)
    for dv, col in ((dv_cat, "J"), (dv_tipo, "B"), (dv_mon, "N")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}201")
    return wb


def libro_arqueo(encabezado: list[tuple], efectivo: list[tuple], pendientes: list[list]) -> Workbook:
    wb = Workbook()
    _hoja_encabezado(wb, encabezado)
    ws = wb.create_sheet("Efectivo")
    _encabezado(ws, 1, ["Denominación", "Cantidad", "Subtotal"], [16, 12, 16])
    for r, (d, c) in enumerate(efectivo, 2):
        ws.cell(row=r, column=1, value=d).number_format = "#,##0"
        ws.cell(row=r, column=2, value=c)
        ws.cell(row=r, column=3, value=f"=A{r}*B{r}").number_format = "#,##0.00"
    total = len(efectivo) + 2
    ws.cell(row=total, column=1, value="TOTAL EFECTIVO").font = Font(bold=True)
    ws.cell(row=total, column=3, value=f"=SUM(C2:C{total - 1})").number_format = "#,##0.00"
    ws = wb.create_sheet("Pendientes")
    _encabezado(ws, 1, ["Tipo", "Fecha", "Clave numérica", "Número", "Beneficiario / proveedor", "Descripción", "Monto"],
                [18, 12, 56, 24, 30, 40, 14])
    dv = DataValidation(type="list", formula1='"Factura por liquidar,Vale,Adelanto"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A2:A100")
    for r in range(2, 101):
        ws.cell(row=r, column=3).number_format = "@"
        ws.cell(row=r, column=4).number_format = "@"
    for r, fila in enumerate(pendientes, 2):
        for c, v in enumerate(fila, 1):
            ws.cell(row=r, column=c, value=v)
    return wb


DENOMINACIONES = [20000, 10000, 5000, 2000, 1000, 500, 100, 50, 25, 10, 5]


def crear_catalogo(ruta: Path, filas: list[list] | None = None) -> Path:
    """Plantilla del catálogo de unidades de negocio (una fila por caja chica)."""
    from .config import COLUMNAS_CATALOGO

    wb = Workbook()
    ws = wb.active
    ws.title = "Cajas"
    _encabezado(ws, 1, COLUMNAS_CATALOGO, [13, 32, 28, 16, 15, 36, 30, 8])
    ws.freeze_panes = "A2"
    for r in range(2, 302):
        ws.cell(row=r, column=4).number_format = "@"
        ws.cell(row=r, column=5).number_format = "#,##0"
    for r, fila in enumerate(filas or [], 2):
        for c, v in enumerate(fila, 1):
            ws.cell(row=r, column=c, value=v)
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H301")
    nota = wb.create_sheet("Instrucciones")
    for i, t in enumerate([
        "Una fila por caja chica (unidad de negocio).",
        "Código caja: el mismo código que el custodio escribe en su liquidación y arqueo (ej. UN-001).",
        "Responsable: custodio autorizado. Si entrega otra persona, el agente lo marca.",
        "Monto del fondo: se usa para cuadrar los arqueos y validar que la liquidación no exceda el fondo.",
        "Aprobadores autorizados: nombres separados por coma. Si se deja vacío, no se valida el aprobador.",
        "Activa: 'No' para cajas cerradas; cualquier liquidación de una caja inactiva se marca como grave.",
    ], 1):
        nota.cell(row=i, column=1, value=t)
    nota.column_dimensions["A"].width = 110
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def crear_plantillas(carpeta: Path, politica: dict) -> list[Path]:
    carpeta.mkdir(parents=True, exist_ok=True)
    liq = libro_liquidacion([
        ("Caja", "CAJA-XXX"), ("Custodio", ""), ("Cédula custodio", ""), ("Número de liquidación", ""),
        ("Fecha de liquidación", ""), ("Periodo desde", ""), ("Periodo hasta", ""), ("Monto solicitado", ""),
    ], [], politica)
    p1 = carpeta / "Plantilla_Liquidacion_CajaChica.xlsx"
    liq.save(p1)
    arq = libro_arqueo([("Caja", "CAJA-XXX"), ("Custodio", ""), ("Fecha del arqueo", ""), ("Realizado por", ""),
                        ("Fondo asignado", "")], [(d, 0) for d in DENOMINACIONES], [])
    p2 = carpeta / "Plantilla_Arqueo_CajaChica.xlsx"
    arq.save(p2)
    p3 = carpeta / "Plantilla_Liquidacion_CajaChica.csv"
    with open(p3, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerow(["Caja", "Custodio", "Número de liquidación", "Fecha de liquidación",
                                "Monto solicitado"] + COLUMNAS_FACTURAS)
    p4 = carpeta / "INSTRUCCIONES.txt"
    inst = politica["institucion"]
    p4.write_text(f"""CÓMO ENTREGAR LA LIQUIDACIÓN Y EL ARQUEO DE CAJA CHICA
=====================================================

1. Descargue la plantilla de esta carpeta (Liquidación o Arqueo). No cambie los nombres de las columnas.
2. Complete una fila por factura. La clave numérica tiene 50 dígitos: péguela tal cual (la columna ya está en formato Texto).
3. TODA factura debe venir a nombre de:
      {inst.get('nombre')}
      Cédula jurídica {inst.get('cedula_juridica')}
   No se aceptan tiquetes electrónicos (no traen los datos de la institución).
4. Guarde el archivo con un nombre como:  Liquidacion_CAJA-ADM_2026-09-18.xlsx  o  Arqueo_CAJA-ADM_2026-09-18.xlsx
5. Súbalo a la carpeta 01_Entrada.
6. En la carpeta 03_Revisiones aparecerá el informe con el resultado, factura por factura.
   Si hay que corregir, corrija el mismo archivo (mismo número de liquidación) y vuelva a subirlo a 01_Entrada.
7. Si el archivo no se pudo leer, aparecerá en 04_Rechazados con una nota explicando por qué.
""", encoding="utf-8")
    return [p1, p2, p3, p4]


# ------------------------------------------------------------------ demostración

def _clave(fecha_: date, cedula: str, numero: int, tipo="01") -> str:
    consecutivo = f"00100001{tipo}{numero:010d}"  # sucursal 001 + terminal 00001 + tipo + número
    return f"506{fecha_:%d%m%y}{solo_digitos(cedula).zfill(12)}{consecutivo}1{random.randint(10**7, 10**8 - 1)}"


CATALOGO_DEMO = [
    ["CAJA-ADM", "Administración", "María Fernández", "1-1111-1111", 500000, "Carlos Rojas", "", "Sí"],
    ["CAJA-OPS", "Operaciones", "Laura Vega", "1-2222-2222", 300000, "Roberto Méndez", "", "Sí"],
    ["CAJA-VEN", "Ventas", "Andrés Quesada", "1-3333-3333", 250000, "Sofía Arias", "", "Sí"],
    ["CAJA-LOG", "Logística", "Diana Brenes", "1-4444-4444", 400000, "Esteban Mora", "", "Sí"],
    ["CAJA-RH", "Recursos Humanos", "Karla Salas", "1-5555-5555", 150000, "Marco Vargas", "", "Sí"],
    ["CAJA-TI", "Tecnología", "Óscar Rojas", "1-6666-6666", 200000, "Marco Vargas", "", "Sí"],
]


def crear_demo(carpeta: Path, politica: dict) -> list[Path]:
    random.seed(7)
    inst = politica["institucion"]
    NOM, CED = inst["nombre"], inst["cedula_juridica"]

    def fac(fecha_, prov, ced_prov, num, desc, cat, total, sol, aut, receptor=NOM, ced_rec=CED, tipo="01", clave=None):
        sub = round(total / 1.13, 2)
        iva = round(total - sub, 2)
        clave = clave or _clave(fecha_, ced_prov, num, tipo)
        return [fecha_, tipo, clave, clave[21:41], prov, ced_prov, receptor, ced_rec, desc, cat, sub, iva, total,
                "CRC", sol, aut]

    d = lambda s: date.fromisoformat(s)
    duplicada = fac(d("2026-09-03"), "Librería El Estudiante S.A.", "3-101-456789", 1520, "Resmas de papel y tóner",
                    "Materiales de oficina", 38500, "Ana Mora", "Carlos Rojas")
    liq1 = [
        duplicada,
        fac(d("2026-09-04"), "Super Central S.A.", "3-101-222333", 88121, "Café, azúcar y galletas para reunión con visitantes",
            "Atención a visitantes", 12450, "Ana Mora", "Carlos Rojas"),
        fac(d("2026-09-05"), "Correos de Costa Rica", "4-000-042151", 3310, "Envío de documentos a oficina regional",
            "Correos y envíos", 4200, "Luis Solís", "Carlos Rojas"),
        list(duplicada),  # la misma factura dos veces en la misma liquidación
        fac(d("2026-09-08"), "Minisúper La Esquina", "1-0987-0654", 776, "Refrescos varios",
            "Alimentación", 6800, "Luis Solís", "Carlos Rojas", tipo="04", receptor="", ced_rec=""),
        fac(d("2026-09-08"), "Ferretería Industrial S.A.", "3-101-777888", 4410, "Candado y cadena bodega",
            "Mantenimiento menor", 15300, "Luis Solís", "Carlos Rojas", receptor="Luis Solís Vargas", ced_rec="1-1234-0567"),
        fac(d("2026-09-09"), "Licorera Nacional", "3-101-999000", 5521, "Vino y cerveza para despedida",
            "Atención a visitantes", 42000, "Ana Mora", "Carlos Rojas"),
        fac(d("2026-09-10"), "Tecno Store S.A.", "3-101-555111", 9912, "Audífonos inalámbricos",
            "Materiales de oficina", 135000, "Ana Mora", "Carlos Rojas"),
    ]
    enc1 = [("Caja", "CAJA-ADM"), ("Custodio", "María Fernández"), ("Cédula custodio", "1-1111-1111"),
            ("Número de liquidación", "ADM-001"), ("Fecha de liquidación", d("2026-09-11")),
            ("Periodo desde", d("2026-09-01")), ("Periodo hasta", d("2026-09-11")),
            ("Monto solicitado", sum(f[12] for f in liq1))]

    liq2 = [
        list(duplicada),  # ya reintegrada en ADM-001
        fac(d("2026-09-14"), "Ferretería Industrial S.A.", "3-101-777888", 4502, "Pintura para pasillo", "Mantenimiento menor",
            48000, "Luis Solís", "Carlos Rojas"),
        fac(d("2026-09-14"), "Ferretería Industrial S.A.", "3-101-777888", 4503, "Brochas y rodillos", "Mantenimiento menor",
            31000, "Luis Solís", "Carlos Rojas"),
        fac(d("2026-09-15"), "Ferretería Industrial S.A.", "3-101-777888", 4519, "Pintura segunda mano", "Mantenimiento menor",
            46500, "Luis Solís", "Carlos Rojas"),
        fac(d("2026-09-16"), "Parqueo Central", "3-101-121212", 802, "Parqueo gestión en Hacienda", "Parqueos y peajes",
            3000, "Carlos Rojas", "Carlos Rojas"),
        fac(d("2026-09-16"), "Soda El Buen Sabor", "1-0555-0444", 1203, "Almuerzo de trabajo con proveedor en restaurante",
            "Alimentación", 18900, "Ana Mora", "Carlos Rojas"),
        fac(d("2026-09-17"), "Super Central S.A.", "3-101-222333", 88590, "Productos de limpieza", "Limpieza",
            22180, "Ana Mora", ""),
    ]
    enc2 = [("Caja", "CAJA-ADM"), ("Custodio", "María Fernández"), ("Cédula custodio", "1-1111-1111"),
            ("Número de liquidación", "ADM-002"), ("Fecha de liquidación", d("2026-09-18")),
            ("Periodo desde", d("2026-09-12")), ("Periodo hasta", d("2026-09-18")),
            ("Monto solicitado", sum(f[12] for f in liq2))]

    rutas = []
    carpeta.mkdir(parents=True, exist_ok=True)
    p = carpeta / "Liquidacion_CAJA-ADM_2026-09-11.xlsx"
    libro_liquidacion(enc1, liq1, politica).save(p)
    rutas.append(p)
    p = carpeta / "Liquidacion_CAJA-ADM_2026-09-18.xlsx"
    libro_liquidacion(enc2, liq2, politica).save(p)
    rutas.append(p)

    # Liquidación en CSV (otra caja), con una diferencia en el monto solicitado
    ops = [
        fac(d("2026-09-15"), "Estación de Servicio La Uruca", "3-101-343434", 70011, "Combustible vehículo institucional",
            "Combustible", 35000, "Pedro Jiménez", "Roberto Méndez"),
        fac(d("2026-09-19"), "Taxi Unidos", "3-101-656565", 3321, "Taxi traslado de equipo", "Transporte",
            7800, "Pedro Jiménez", "Roberto Méndez"),
        fac(d("2026-08-10"), "Fotocopiadora Rápida", "1-0777-0888", 2210, "Copias de expedientes", "Fotocopias e impresión",
            5650, "Pedro Jiménez", "Roberto Méndez"),
        fac(d("2026-09-20"), "Pedro Jiménez Arias", "1-0999-0111", 15, "Reparación menor de silla", "Mantenimiento menor",
            20000, "Pedro Jiménez", "Roberto Méndez"),
        fac(d("2026-09-21"), "Imprenta Gráfica S.A.", "3-101-898989", 6610, "Impresión de rótulos", "Publicidad",
            27400, "Pedro Jiménez", "Roberto Méndez"),
    ]
    ops[1][11] = 1900  # IVA mal digitado
    p = carpeta / "Liquidacion_CAJA-OPS_2026-09-22.csv"
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Caja", "Custodio", "Número de liquidación", "Fecha de liquidación", "Monto solicitado"] + COLUMNAS_FACTURAS)
        for fila in ops:
            w.writerow(["CAJA-OPS", "Laura Vega", "OPS-010", "22/09/2026", "100000"] +
                       [x.strftime("%d/%m/%Y") if isinstance(x, date) else x for x in fila])
    rutas.append(p)

    # Arqueo con faltante, un vale vencido y una factura pendiente que ya fue reintegrada
    efectivo = [(20000, 15), (10000, 8), (5000, 4), (2000, 5), (1000, 1), (500, 1), (100, 2)]
    pendientes = [
        ["Factura por liquidar", d("2026-09-19"), liq2[4][2], liq2[4][3], "Parqueo Central", "Parqueo", 3000],
        ["Factura por liquidar", d("2026-09-03"), duplicada[2], duplicada[3], "Librería El Estudiante S.A.",
         "Resmas de papel y tóner", 38500],
        ["Vale", d("2026-09-10"), "", "V-118", "Luis Solís", "Compra de repuestos (pendiente factura)", 25000],
        ["Factura por liquidar", d("2026-09-22"), "", "771", "Super Central S.A.", "Agua embotellada", 9800],
    ]
    p = carpeta / "Arqueo_CAJA-ADM_2026-09-23.xlsx"
    libro_arqueo([("Caja", "CAJA-ADM"), ("Custodio", "María Fernández"), ("Fecha del arqueo", d("2026-09-23")),
                  ("Realizado por", "Jorge Castro (Auditoría)"), ("Fondo asignado", 500000)], efectivo, pendientes).save(p)
    rutas.append(p)
    return rutas
