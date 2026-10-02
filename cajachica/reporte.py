"""Informe de revisión en Excel que se devuelve al repositorio por cada archivo procesado."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .verificacion import ACCIONES, AVISO_ART_23

AZUL = "1F3A5F"
COLORES_ESTADO = {
    "RECHAZADA": "F8D7DA", "OBSERVADA": "FFF3CD", "OK": "D4EDDA",
    "ALTA": "F8D7DA", "MEDIA": "FFF3CD", "BAJA": "E2E3E5",
    "CUMPLE": "D4EDDA", "NO CUMPLE": "F8D7DA", "NO SE PUEDE VERIFICAR": "FFF3CD", "NO APLICA": "E2E3E5",
}
BORDE = Border(bottom=Side(style="thin", color="DDDDDD"))
FMT_MONTO = '"₡"#,##0.00'
FMT_FECHA = "DD/MM/YYYY"


def _titulo(ws, texto, subtitulo=""):
    ws["A1"] = texto
    ws["A1"].font = Font(size=15, bold=True, color=AZUL)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(size=10, italic=True, color="666666")


def _tabla(ws, fila_inicio, encabezados, filas, anchos=None, col_estado=None, montos=(), fechas=()):
    for c, h in enumerate(encabezados, 1):
        celda = ws.cell(row=fila_inicio, column=c, value=h)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for r, fila in enumerate(filas, fila_inicio + 1):
        for c, v in enumerate(fila, 1):
            celda = ws.cell(row=r, column=c, value=v)
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            celda.border = BORDE
            if c in montos:
                celda.number_format = FMT_MONTO
            if c in fechas:
                celda.number_format = FMT_FECHA
        if col_estado:
            estado = fila[col_estado - 1]
            if estado in COLORES_ESTADO:
                ws.cell(row=r, column=col_estado).fill = PatternFill("solid", fgColor=COLORES_ESTADO[estado])
                ws.cell(row=r, column=col_estado).font = Font(bold=True)
    for i, ancho in enumerate(anchos or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{fila_inicio}:{get_column_letter(len(encabezados))}{fila_inicio + len(filas)}"


def _pares(ws, fila, pares):
    for etiqueta, valor, *fmt in pares:
        ws.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = ws.cell(row=fila, column=2, value=valor)
        if fmt:
            celda.number_format = fmt[0]
        fila += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 48
    return fila


def _hoja_hallazgos(ws, filas):
    """Tabla del reglamento: n.º, documento, hallazgo, artículo, gravedad sugerida, acción recomendada."""
    orden = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    filas = sorted(filas, key=lambda x: orden[x[0]["severidad"]])
    datos = [[i, h["severidad"], doc, f"{h['titulo']}: {h['detalle']}", h.get("articulo", ""),
              h.get("gravedad", "") or "—", ACCIONES.get(h["regla"], ""), monto]
             for i, (h, doc, monto) in enumerate(filas, 1)]
    _tabla(ws, 1, ["N.º", "Prioridad", "Documento", "Hallazgo", "Artículo", "Gravedad sugerida", "Acción recomendada",
                   "Monto ₡"], datos, [5, 10, 40, 70, 30, 22, 44, 13], col_estado=2, montos=(8,))


def _hojas_verificacion(wb, documentos):
    ws = wb.create_sheet("Lista de verificación")
    filas = [[doc, it["grupo"], it["control"], it["articulo"], it["estado"], it["detalle"]]
             for doc, r in documentos for it in r.get("verificacion", [])]
    _tabla(ws, 1, ["Documento", "Área", "Control", "Artículo", "Estado", "Detalle"], filas,
           [16, 24, 52, 22, 22, 80], col_estado=5)
    cumplen = [[doc, it["control"], it["articulo"]] for doc, r in documentos for it in r.get("verificacion", [])
               if it["estado"] == "CUMPLE"]
    ws = wb.create_sheet("Puntos que cumplen")
    _tabla(ws, 1, ["Documento", "Control", "Artículo"], cumplen, [16, 60, 26])
    ws = wb.create_sheet("Documentos pendientes")
    pend = [[doc, p] for doc, r in documentos for p in r.get("pendientes", [])]
    _tabla(ws, 1, ["Documento", "Qué falta para completar la revisión"], pend, [16, 120])


def informe_liquidaciones(resultados: list[dict], archivo: str, destino: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "Revisión de liquidación de caja chica",
            f"Archivo: {archivo} · Revisado: {datetime.now():%d/%m/%Y %H:%M}")
    fila = 4
    for res in resultados:
        enc, r = res["encabezado"], res["resumen"]
        fila = _pares(ws, fila, [
            ("Caja (FUC)", enc["caja"]), ("Persona encargada", enc["custodio"]), ("N° de liquidación", enc["numero"]),
            ("Fecha de liquidación", enc["fecha_liquidacion"], FMT_FECHA),
            ("ESTADO", r["estado"]),
            ("Facturas presentadas", r["n_facturas"]),
            ("Total presentado", r["total"], FMT_MONTO),
            ("Monto solicitado", enc["monto_solicitado"], FMT_MONTO),
            ("Facturas rechazadas", r["n_rechazadas"]), ("Monto rechazado", r["monto_rechazado"], FMT_MONTO),
            ("Facturas observadas", r["n_observadas"]), ("Monto observado", r["monto_observado"], FMT_MONTO),
            ("Monto que podría reintegrarse", r["monto_aprobable"], FMT_MONTO),
        ])
        estado_celda = ws.cell(row=fila - 9, column=2)
        estado_celda.font = Font(bold=True, color="FFFFFF")
        estado_celda.fill = PatternFill("solid", fgColor={"CON RECHAZOS": "C0392B", "CON OBSERVACIONES": "D68910"}
                                        .get(r["estado"], "1E8449"))
        fila += 1
    ws.cell(row=fila + 1, column=1, value="Cómo leer este informe").font = Font(bold=True, color=AZUL)
    for i, t in enumerate([
        "RECHAZADA: la factura no debe reintegrarse hasta corregir el hallazgo de severidad ALTA.",
        "OBSERVADA: requiere justificación o corrección antes de aprobar.",
        "OK: sin hallazgos que impidan el reintegro (puede tener alertas informativas).",
        "Corrija el archivo y vuelva a cargarlo en 01_Entrada con el mismo N° de liquidación: reemplaza la versión anterior.",
    ], fila + 2):
        ws.cell(row=i, column=1, value=t)

    ws.cell(row=fila + 7, column=1, value=AVISO_ART_23).font = Font(italic=True, color="666666")

    _hojas_verificacion(wb, [(r["encabezado"]["numero"], r) for r in resultados])

    ws = wb.create_sheet("Factura por factura")
    filas = []
    for res in resultados:
        for f in res["facturas"]:
            obs = "\n".join(f"[{h['severidad']}] {h['titulo']}: {h['detalle']}" for h in f["hallazgos"])
            filas.append([res["encabezado"]["numero"], f["fila"], f["estado"], f["fecha"], f["proveedor"],
                          f["cedula_proveedor"], f["consecutivo"], f["receptor"], f["cedula_receptor"],
                          f["descripcion"], f["categoria"], f["total_crc"], obs or "Sin observaciones"])
    _tabla(ws, 1, ["Liquidación", "Fila", "Estado", "Fecha", "Proveedor", "Cédula prov.", "N° factura",
                   "Receptor", "Cédula receptor", "Descripción", "Categoría", "Total ₡", "Observaciones"],
           filas, [12, 6, 12, 11, 26, 14, 22, 26, 14, 34, 18, 13, 80], col_estado=3, montos=(12,), fechas=(4,))

    ws = wb.create_sheet("Hallazgos")
    filas = []
    for res in resultados:
        n = res["encabezado"]["numero"]
        for h in res["hallazgos_generales"]:
            filas.append([h, f"Liquidación {n}", h["monto"]])
        for f in res["facturas"]:
            doc = f"Liq. {n} · fila {f['fila']} · {f['proveedor']} · factura {f['consecutivo'] or '—'}"
            for h in f["hallazgos"]:
                filas.append([h, doc, h["monto"] if h["monto"] is not None else f["total_crc"]])
    _hoja_hallazgos(ws, filas)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def informe_arqueos(resultados: list[dict], archivo: str, destino: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "Revisión de arqueo de caja chica", f"Archivo: {archivo} · Revisado: {datetime.now():%d/%m/%Y %H:%M}")
    fila = 4
    for a in resultados:
        fila = _pares(ws, fila, [
            ("Caja", a["caja"]), ("Custodio", a["custodio"]), ("Fecha del arqueo", a["fecha"], FMT_FECHA),
            ("Realizado por", a["realizado_por"]), ("ESTADO", a["estado"]),
            ("Fondo", a["fondo"], FMT_MONTO), ("Efectivo contado", a["efectivo"], FMT_MONTO),
            ("Facturas pendientes de liquidar", a["pendientes_facturas"], FMT_MONTO),
            ("Vales / adelantos", a["vales"], FMT_MONTO),
            ("Diferencia (+ sobrante / − faltante)", a["diferencia"], FMT_MONTO),
        ]) + 1
    ws.cell(row=fila + 1, column=1, value=AVISO_ART_23).font = Font(italic=True, color="666666")
    _hojas_verificacion(wb, [(f"Arqueo {a['caja']}", a) for a in resultados])
    ws = wb.create_sheet("Hallazgos")
    _hoja_hallazgos(ws, [[h, f"Arqueo {a['caja']} del {a['fecha']:%d/%m/%Y}" if a["fecha"] else f"Arqueo {a['caja']}",
                          h["monto"]] for a in resultados for h in a["hallazgos"]])
    ws = wb.create_sheet("Pendientes")
    filas = [[a["caja"], p["tipo"], p["fecha"], p["numero"], p["beneficiario"], p["descripcion"], p["monto"]]
             for a in resultados for p in a["detalle_pendientes"]]
    _tabla(ws, 1, ["Caja", "Tipo", "Fecha", "Número", "Beneficiario / proveedor", "Descripción", "Monto ₡"],
           filas, [12, 14, 12, 24, 30, 40, 14], montos=(7,), fechas=(3,))
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
