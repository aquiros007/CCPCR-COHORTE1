"""Comprobantes electrónicos que entregan los proveedores: PDF + XML de la factura + respuesta de Hacienda.

Los tres archivos llegan por separado a 07_Proveedores/{01_PDF, 02_XML, 03_Respuesta_Hacienda}. Se unen por la
clave numérica de 50 dígitos (no importa cómo se llamen los archivos), se cruzan entre sí y el resultado queda
para la sección de aprobación del dashboard.
"""
import json
import re
import shutil
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .config import RUTA_CONFIG, ruta_repositorio
from .util import TIPOS_DOCUMENTO, descomponer_clave, fecha, monto, normalizar, solo_digitos, texto

RUTA_PROVEEDORES = RUTA_CONFIG / "Proveedores.xlsx"
COLUMNAS_PROVEEDOR = ["Cédula", "Razón social", "Nombre comercial", "Correo", "Teléfono", "Estado",
                      "Fecha de registro", "Observaciones"]
CARPETAS = {"registro": "00_Registro", "pdf": "01_PDF", "xml": "02_XML", "respuesta": "03_Respuesta_Hacienda",
            "conciliados": "04_Conciliados", "no_legibles": "05_No_legibles"}
TIPOS_FACTURA = {"FacturaElectronica", "TiqueteElectronico", "NotaCreditoElectronica", "NotaDebitoElectronica",
                 "FacturaElectronicaCompra", "FacturaElectronicaExportacion"}
ESTADO_HACIENDA = {"1": "Aceptado", "2": "Aceptado parcialmente", "3": "Rechazado"}

CUMPLE, NO_CUMPLE, NO_VERIFICABLE, NO_APLICA = "CUMPLE", "NO CUMPLE", "NO SE PUEDE VERIFICAR", "NO APLICA"


def carpetas_proveedores(politica: dict) -> dict:
    base = ruta_repositorio(politica) / "07_Proveedores"
    rutas = {k: base / v for k, v in CARPETAS.items()}
    for r in rutas.values():
        r.mkdir(parents=True, exist_ok=True)
    return rutas


# ------------------------------------------------------------------ lectura de archivos

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _buscar(raiz, ruta: str):
    """Busca por nombres locales (ignora el namespace, que cambia entre versiones 4.3 y 4.4)."""
    nodos = [raiz]
    for parte in ruta.split("/"):
        nodos = [h for n in nodos for h in list(n) if _local(h.tag) == parte]
        if not nodos:
            return None
    return nodos[0]


def _txt(raiz, ruta):
    n = _buscar(raiz, ruta)
    return (n.text or "").strip() if n is not None else ""


def leer_xml(ruta: Path) -> dict:
    return leer_xml_bytes(ruta.read_bytes())


def leer_xml_bytes(contenido: bytes) -> dict:
    """Lee un XML de Hacienda desde bytes (archivo en disco o recibido por la web).
    Rechaza DTD y entidades para evitar ataques de expansión de entidades."""
    if b"<!DOCTYPE" in contenido[:2048].upper() or b"<!ENTITY" in contenido.upper():
        return {"tipo": "ilegible", "error": "El XML trae declaraciones DTD/ENTITY, que no se aceptan."}
    try:
        raiz = ET.fromstring(contenido)
    except ET.ParseError as e:
        return {"tipo": "ilegible", "error": f"XML mal formado: {e}"}
    tipo = _local(raiz.tag)
    if tipo == "MensajeHacienda":
        return {
            "tipo": "respuesta",
            "clave": solo_digitos(_txt(raiz, "Clave")),
            "emisor_nombre": _txt(raiz, "NombreEmisor"),
            "emisor_cedula": solo_digitos(_txt(raiz, "NumeroCedulaEmisor")),
            "receptor_nombre": _txt(raiz, "NombreReceptor"),
            "receptor_cedula": solo_digitos(_txt(raiz, "NumeroCedulaReceptor")),
            "mensaje": _txt(raiz, "Mensaje"),
            "detalle": _txt(raiz, "DetalleMensaje"),
            "impuesto": monto(_txt(raiz, "MontoTotalImpuesto")),
            "total": monto(_txt(raiz, "TotalFactura")),
        }
    if tipo == "MensajeReceptor":
        return {"tipo": "mensaje_receptor", "clave": solo_digitos(_txt(raiz, "Clave")),
                "error": "Es la aceptación del receptor (MensajeReceptor), no la respuesta de Hacienda."}
    if tipo in TIPOS_FACTURA:
        nodo_detalle = _buscar(raiz, "DetalleServicio")
        detalle = [_txt(l, "Detalle") for l in (list(nodo_detalle) if nodo_detalle is not None else [])]
        return {
            "tipo": "factura",
            "documento": tipo,
            "clave": solo_digitos(_txt(raiz, "Clave")),
            "consecutivo": solo_digitos(_txt(raiz, "NumeroConsecutivo")),
            "fecha": fecha(_txt(raiz, "FechaEmision")[:10]),
            "emisor_nombre": _txt(raiz, "Emisor/Nombre"),
            "emisor_cedula": solo_digitos(_txt(raiz, "Emisor/Identificacion/Numero")),
            "emisor_correo": _txt(raiz, "Emisor/CorreoElectronico"),
            "receptor_nombre": _txt(raiz, "Receptor/Nombre"),
            "receptor_cedula": solo_digitos(_txt(raiz, "Receptor/Identificacion/Numero")),
            "moneda": _txt(raiz, "ResumenFactura/CodigoTipoMoneda/CodigoMoneda") or "CRC",
            "impuesto": monto(_txt(raiz, "ResumenFactura/TotalImpuesto")) or 0.0,
            "total": monto(_txt(raiz, "ResumenFactura/TotalComprobante")),
            "descripcion": "; ".join(d for d in detalle if d)[:500],
            "firmado": any(_local(e.tag) == "Signature" for e in raiz.iter()),
        }
    return {"tipo": "ilegible", "error": f"No es un comprobante de Hacienda (elemento raíz '{tipo}')."}


def leer_pdf(ruta: Path) -> dict:
    return leer_pdf_bytes(ruta.read_bytes(), ruta.stem)


def leer_pdf_bytes(contenido: bytes, nombre: str = "") -> dict:
    try:
        import io
        from pypdf import PdfReader
        texto_pdf = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(contenido)).pages[:10])
    except Exception as e:  # PDF dañado o protegido
        return {"tipo": "pdf", "texto": "", "error": f"No se pudo leer el PDF: {e}"}
    digitos = re.sub(r"\D", "", texto_pdf)
    claves = sorted(set(re.findall(r"506\d{47}", digitos)))
    # Clave también puede venir en el nombre del archivo
    claves += [c for c in re.findall(r"506\d{47}", re.sub(r"\D", "", nombre)) if c not in claves]
    return {"tipo": "pdf", "texto": texto_pdf, "digitos": digitos, "claves": claves,
            "escaneado": len(texto_pdf.strip()) < 30}


def _montos_en_texto(texto_pdf: str) -> set:
    valores = set()
    for m in re.findall(r"\d[\d.,\s]*\d", texto_pdf):
        v = monto(m.replace(" ", ""))
        if v is not None:
            valores.add(round(v, 2))
    return valores


# ------------------------------------------------------------------ registro de proveedores

def cargar_proveedores(ruta: Path = RUTA_PROVEEDORES) -> dict:
    if not ruta.exists():
        return {}
    df = pd.read_excel(ruta, dtype=object)
    col = {normalizar(c): c for c in df.columns}
    v = lambda fila, n: fila[col[normalizar(n)]] if normalizar(n) in col else None
    return {solo_digitos(v(f, "Cédula")): {
        "cedula": solo_digitos(v(f, "Cédula")), "razon_social": texto(v(f, "Razón social")),
        "nombre_comercial": texto(v(f, "Nombre comercial")), "correo": texto(v(f, "Correo")),
        "telefono": texto(v(f, "Teléfono")), "estado": texto(v(f, "Estado")) or "Pendiente",
        "fecha_registro": texto(v(f, "Fecha de registro")), "observaciones": texto(v(f, "Observaciones")),
    } for _, f in df.iterrows() if solo_digitos(v(f, "Cédula"))}


def guardar_proveedores(proveedores: dict, ruta: Path = RUTA_PROVEEDORES) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Proveedores"
    for c, h in enumerate(COLUMNAS_PROVEEDOR, 1):
        celda = ws.cell(row=1, column=c, value=h)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
    for r, p in enumerate(sorted(proveedores.values(), key=lambda p: p["razon_social"]), 2):
        for c, k in enumerate(["cedula", "razon_social", "nombre_comercial", "correo", "telefono", "estado",
                               "fecha_registro", "observaciones"], 1):
            ws.cell(row=r, column=c, value=p.get(k, "")).number_format = "@"
    for letra, ancho in zip("ABCDEFGH", [16, 36, 26, 30, 14, 12, 14, 40]):
        ws.column_dimensions[letra].width = ancho
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def procesar_registros(carpeta: Path, proveedores: dict) -> list[dict]:
    """Formularios de registro (Excel) que los proveedores dejan en 00_Registro."""
    nuevos = []
    for ruta in sorted(carpeta.glob("*.xlsx")):
        if ruta.name.startswith("~$"):
            continue
        df = pd.read_excel(ruta, header=None, dtype=object)
        datos = {}
        for _, fila in df.iterrows():
            vals = [x for x in fila.tolist() if texto(x)]
            if len(vals) >= 2:
                datos[normalizar(vals[0])] = vals[1]
        cedula = solo_digitos(datos.get("CEDULA") or datos.get("CEDULA JURIDICA O FISICA"))
        if not cedula:
            nuevos.append({"archivo": ruta.name, "resultado": "Sin cédula: no se registró"})
            continue
        previo = proveedores.get(cedula, {})
        proveedores[cedula] = {
            "cedula": cedula, "razon_social": texto(datos.get("RAZON SOCIAL")) or previo.get("razon_social", ""),
            "nombre_comercial": texto(datos.get("NOMBRE COMERCIAL")), "correo": texto(datos.get("CORREO")),
            "telefono": texto(datos.get("TELEFONO")),
            "estado": previo.get("estado") or "Pendiente",
            "fecha_registro": previo.get("fecha_registro") or date.today().isoformat(),
            "observaciones": previo.get("observaciones", ""),
        }
        procesado = carpeta / "procesados"
        procesado.mkdir(exist_ok=True)
        shutil.move(str(ruta), procesado / ruta.name)
        nuevos.append({"archivo": ruta.name, "cedula": cedula,
                       "resultado": "Actualizado" if previo else "Registrado (pendiente de aprobación)"})
    return nuevos


# ------------------------------------------------------------------ cruce de los tres archivos

def _control(nombre, estado, detalle=""):
    return {"control": nombre, "estado": estado, "detalle": detalle}


def cruzar(clave: str, fac: dict | None, resp: dict | None, pdf: dict | None, politica: dict,
           proveedores: dict, duplicada: bool) -> list[dict]:
    inst = solo_digitos(politica["institucion"]["cedula_juridica"])
    tol = float((politica.get("proveedores") or {}).get("tolerancia_monto", 1))
    exigir_aprobado = (politica.get("proveedores") or {}).get("exigir_proveedor_aprobado", True)
    c = []
    faltan = [n for n, x in (("PDF", pdf), ("XML de la factura", fac), ("respuesta de Hacienda", resp)) if not x]
    c.append(_control("Archivos completos", NO_CUMPLE if faltan else CUMPLE,
                      ("Falta: " + ", ".join(faltan)) if faltan else "PDF, XML y respuesta de Hacienda recibidos."))

    partes = descomponer_clave(clave)
    tipo = partes["tipo_documento"] if partes else ""
    aceptados = [str(t).zfill(2) for t in politica["comprobantes"].get("tipos_documento_aceptados", ["01"])]
    c.append(_control("Tipo de comprobante aceptado",
                      NO_VERIFICABLE if not partes else CUMPLE if tipo in aceptados else NO_CUMPLE,
                      TIPOS_DOCUMENTO.get(tipo, f"Tipo {tipo}") if partes else "Clave inválida."))

    if resp:
        estado = ESTADO_HACIENDA.get(resp["mensaje"], f"Código {resp['mensaje']}")
        c.append(_control("Hacienda aceptó el comprobante", CUMPLE if resp["mensaje"] == "1" else NO_CUMPLE,
                          estado + (f": {resp['detalle'][:200]}" if resp["mensaje"] != "1" and resp["detalle"] else "")))
    else:
        c.append(_control("Hacienda aceptó el comprobante", NO_VERIFICABLE, "Falta la respuesta de Hacienda."))

    if fac and resp:
        c.append(_control("Emisor coincide (XML y respuesta)",
                          CUMPLE if fac["emisor_cedula"] == resp["emisor_cedula"] else NO_CUMPLE,
                          f"XML {fac['emisor_cedula']} · respuesta {resp['emisor_cedula']}"))
        dif_t = abs((fac["total"] or 0) - (resp["total"] or 0))
        c.append(_control("Total coincide (XML y respuesta)", CUMPLE if dif_t <= tol else NO_CUMPLE,
                          f"XML {fac['total']:,.2f} · respuesta {resp['total'] or 0:,.2f}"))
        dif_i = abs((fac["impuesto"] or 0) - (resp["impuesto"] or 0))
        c.append(_control("IVA coincide (XML y respuesta)", CUMPLE if dif_i <= tol else NO_CUMPLE,
                          f"XML {fac['impuesto']:,.2f} · respuesta {resp['impuesto'] or 0:,.2f}"))
    else:
        for n in ("Emisor coincide (XML y respuesta)", "Total coincide (XML y respuesta)", "IVA coincide (XML y respuesta)"):
            c.append(_control(n, NO_VERIFICABLE, "Falta el XML o la respuesta de Hacienda."))

    receptores = [x["receptor_cedula"] for x in (fac, resp) if x]
    if receptores:
        ok = all(r == inst for r in receptores)
        c.append(_control("A nombre de la Universidad Nacional", CUMPLE if ok else NO_CUMPLE,
                          "Cédula del receptor: " + " / ".join(r or "(vacía)" for r in receptores)))
    else:
        c.append(_control("A nombre de la Universidad Nacional", NO_VERIFICABLE, "Sin XML ni respuesta."))

    if fac:
        c.append(_control("XML con firma digital", CUMPLE if fac["firmado"] else NO_CUMPLE,
                          "Contiene la firma del emisor (no se validó criptográficamente)." if fac["firmado"]
                          else "El XML no trae firma digital."))
        if partes and fac["fecha"] and partes["fecha"] and partes["fecha"] != fac["fecha"]:
            c.append(_control("Fecha de emisión coincide con la clave", NO_CUMPLE,
                              f"Clave {partes['fecha']:%d/%m/%Y} · XML {fac['fecha']:%d/%m/%Y}"))
        else:
            c.append(_control("Fecha de emisión coincide con la clave", CUMPLE if partes else NO_VERIFICABLE))

    if not pdf:
        for n in ("El PDF corresponde a la factura", "Total del PDF coincide con el XML", "Cédula del emisor en el PDF"):
            c.append(_control(n, NO_VERIFICABLE, "Falta el PDF."))
    elif pdf.get("escaneado"):
        for n in ("El PDF corresponde a la factura", "Total del PDF coincide con el XML", "Cédula del emisor en el PDF"):
            c.append(_control(n, NO_VERIFICABLE, "PDF escaneado o sin texto: requiere revisión visual."))
    else:
        c.append(_control("El PDF corresponde a la factura", CUMPLE if clave in pdf["digitos"] else NO_CUMPLE,
                          "La clave del XML aparece en el PDF." if clave in pdf["digitos"]
                          else "La clave del XML no aparece en el PDF."))
        if fac and fac["total"] is not None:
            montos = _montos_en_texto(pdf["texto"])
            ok = any(abs(m - fac["total"]) <= tol for m in montos)
            c.append(_control("Total del PDF coincide con el XML", CUMPLE if ok else NO_CUMPLE,
                              f"Total del XML {fac['total']:,.2f} {'presente' if ok else 'no aparece'} en el PDF."))
            ced = fac["emisor_cedula"]
            c.append(_control("Cédula del emisor en el PDF", CUMPLE if ced and ced in pdf["digitos"] else NO_CUMPLE,
                              f"Cédula {ced}"))
        else:
            c.append(_control("Total del PDF coincide con el XML", NO_VERIFICABLE, "Falta el XML."))
            c.append(_control("Cédula del emisor en el PDF", NO_VERIFICABLE, "Falta el XML."))

    emisor = (fac or resp or {}).get("emisor_cedula", "")
    prov = proveedores.get(emisor)
    if not emisor:
        c.append(_control("Proveedor registrado", NO_VERIFICABLE, "Sin XML ni respuesta para identificar al emisor."))
    elif not prov:
        c.append(_control("Proveedor registrado", NO_CUMPLE, f"La cédula {emisor} no está en el registro de proveedores."))
    elif exigir_aprobado and normalizar(prov["estado"]) != "APROBADO":
        c.append(_control("Proveedor registrado", NO_CUMPLE, f"Registro en estado '{prov['estado']}'."))
    else:
        c.append(_control("Proveedor registrado", CUMPLE, prov["razon_social"]))

    c.append(_control("Comprobante no presentado antes", NO_CUMPLE if duplicada else CUMPLE,
                      "Esta clave ya se había recibido en otro archivo." if duplicada else ""))
    return c


def estado_general(controles: list[dict], vencido: bool) -> str:
    estados = {x["estado"] for x in controles}
    incompleto = any(x["control"] == "Archivos completos" and x["estado"] == NO_CUMPLE for x in controles)
    if incompleto and not vencido:
        return "INCOMPLETO"
    if NO_CUMPLE in estados:
        return "CON DIFERENCIAS"
    if NO_VERIFICABLE in estados:
        return "REVISIÓN MANUAL"
    return "LISTO PARA APROBACIÓN"


# ------------------------------------------------------------------ proceso

def procesar_comprobantes(politica: dict, db) -> dict:
    rutas = carpetas_proveedores(politica)
    proveedores = cargar_proveedores()
    bitacora = {"registros": procesar_registros(rutas["registro"], proveedores), "comprobantes": [], "no_legibles": [],
                "sin_asociar": []}

    grupos: dict[str, dict] = {}
    def grupo(clave):
        return grupos.setdefault(clave, {"pdf": None, "factura": None, "respuesta": None, "archivos": {}})

    for ruta in sorted(rutas["xml"].glob("*")) + sorted(rutas["respuesta"].glob("*")):
        if ruta.is_dir() or ruta.name.startswith((".", "~$")):
            continue
        datos = leer_xml(ruta) if ruta.suffix.lower() == ".xml" else {"tipo": "ilegible", "error": "No es un archivo XML."}
        if datos["tipo"] in ("factura", "respuesta") and len(datos["clave"]) == 50:
            g = grupo(datos["clave"])
            g[datos["tipo"]] = datos
            g["archivos"][datos["tipo"]] = ruta
        else:
            destino = rutas["no_legibles"] / ruta.name
            shutil.move(str(ruta), destino)
            destino.with_suffix(".LEAME.txt").write_text(
                f"El archivo '{ruta.name}' no se pudo usar.\nMotivo: {datos.get('error', 'clave inválida')}\n"
                "Cargue el XML de la factura en 02_XML y el XML de respuesta de Hacienda en 03_Respuesta_Hacienda.\n",
                encoding="utf-8")
            bitacora["no_legibles"].append({"archivo": ruta.name, "motivo": datos.get("error", "Clave inválida")})

    pdfs_sin_clave = []
    for ruta in sorted(rutas["pdf"].glob("*")):
        if ruta.is_dir() or ruta.name.startswith((".", "~$")):
            continue
        if ruta.suffix.lower() != ".pdf":
            continue
        datos = leer_pdf(ruta)
        clave = next((c for c in datos.get("claves", []) if c in grupos), None) or \
            (datos["claves"][0] if datos.get("claves") else None)
        if not clave:
            # PDF escaneado: se intenta por el consecutivo de 20 dígitos en el nombre del archivo
            nombre = re.sub(r"\D", "", ruta.stem)
            clave = next((k for k, g in grupos.items() if g["factura"] and g["factura"]["consecutivo"]
                          and g["factura"]["consecutivo"] in nombre), None)
        if clave:
            g = grupo(clave)
            g["pdf"] = datos
            g["archivos"]["pdf"] = ruta
        else:
            pdfs_sin_clave.append(ruta)

    # Combinar con lo que ya estaba registrado de corridas anteriores (archivos ya conciliados o en espera)
    previos = {r["clave"]: r for r in db.consultar("SELECT * FROM comprobantes")}
    dias_espera = int((politica.get("proveedores") or {}).get("dias_espera_completar", 5))
    hoy = date.today()
    for clave, g in grupos.items():
        previo = previos.get(clave)
        archivos_previos = json.loads(previo["archivos"]) if previo else {}
        # La clave ya estaba completa y vuelve a llegar algún archivo: se presenta de nuevo
        duplicada = bool(previo and all(archivos_previos.get(t) for t in ("pdf", "factura", "respuesta"))
                         and any(str(p) not in archivos_previos.values() for p in g["archivos"].values()))
        for tipo in ("pdf", "factura", "respuesta"):
            if g[tipo] is None and archivos_previos.get(tipo):
                p = Path(archivos_previos[tipo])
                if p.exists():
                    g[tipo] = leer_pdf(p) if tipo == "pdf" else leer_xml(p)
                    g["archivos"][tipo] = p
        recibido = fecha(previo["fecha_recepcion"]) if previo else hoy
        vencido = (hoy - recibido).days >= dias_espera
        controles = cruzar(clave, g["factura"], g["respuesta"], g["pdf"], politica, proveedores, duplicada)
        estado = estado_general(controles, vencido)

        completo = all(g[t] for t in ("pdf", "factura", "respuesta"))
        if completo:  # el juego completo se archiva junto
            destino = rutas["conciliados"] / datetime.now().strftime("%Y-%m") / clave
            destino.mkdir(parents=True, exist_ok=True)
            for tipo, p in list(g["archivos"].items()):
                if p.parent != destino:
                    nuevo = destino / p.name
                    if nuevo.exists():
                        nuevo = destino / f"{p.stem}_{datetime.now():%Y%m%d%H%M%S}{p.suffix}"
                    shutil.move(str(p), nuevo)
                    g["archivos"][tipo] = nuevo

        fac, resp = g["factura"] or {}, g["respuesta"] or {}
        emisor = fac.get("emisor_cedula") or resp.get("emisor_cedula", "")
        if emisor and emisor not in proveedores:  # se registra solo, pendiente de aprobación
            proveedores[emisor] = {"cedula": emisor, "razon_social": fac.get("emisor_nombre") or resp.get("emisor_nombre", ""),
                                   "nombre_comercial": "", "correo": fac.get("emisor_correo", ""), "telefono": "",
                                   "estado": "Pendiente", "fecha_registro": hoy.isoformat(),
                                   "observaciones": "Registrado automáticamente desde su primer XML."}
        fila = {
            "clave": clave, "tipo_doc": (descomponer_clave(clave) or {}).get("tipo_documento", ""),
            "consecutivo": fac.get("consecutivo", ""), "fecha": fac.get("fecha"),
            "emisor_cedula": emisor, "emisor_nombre": fac.get("emisor_nombre") or resp.get("emisor_nombre", ""),
            "receptor_cedula": fac.get("receptor_cedula") or resp.get("receptor_cedula", ""),
            "receptor_nombre": fac.get("receptor_nombre") or resp.get("receptor_nombre", ""),
            "moneda": fac.get("moneda", "CRC"), "total": fac.get("total", resp.get("total")),
            "impuesto": fac.get("impuesto", resp.get("impuesto")), "descripcion": fac.get("descripcion", ""),
            "estado_hacienda": ESTADO_HACIENDA.get(resp.get("mensaje", ""), "Sin respuesta"),
            "archivos": json.dumps({k: str(v) for k, v in g["archivos"].items()}, ensure_ascii=False),
            "controles": json.dumps(controles, ensure_ascii=False, default=str), "estado": estado,
            "fecha_recepcion": previo["fecha_recepcion"] if previo else hoy.isoformat(),
            "fecha_actualizacion": datetime.now().isoformat(timespec="seconds"),
        }
        if previo:
            db.con.execute("DELETE FROM comprobantes WHERE clave=?", (clave,))
        db.insertar("comprobantes", fila)
        bitacora["comprobantes"].append({"clave": clave, "emisor": fila["emisor_nombre"], "total": fila["total"],
                                         "estado": estado,
                                         "diferencias": [x["control"] for x in controles if x["estado"] == NO_CUMPLE]})

    # Juegos incompletos de corridas anteriores que no recibieron nada nuevo: se re-evalúa el vencimiento
    for clave, previo in previos.items():
        if clave in grupos or previo["estado"] != "INCOMPLETO":
            continue
        if (hoy - fecha(previo["fecha_recepcion"])).days >= dias_espera:
            controles = json.loads(previo["controles"])
            db.con.execute("UPDATE comprobantes SET estado=?, fecha_actualizacion=? WHERE clave=?",
                           (estado_general(controles, True), datetime.now().isoformat(timespec="seconds"), clave))

    for ruta in pdfs_sin_clave:
        bitacora["sin_asociar"].append({"archivo": ruta.name, "motivo":
                                        "No se encontró la clave en el PDF ni un XML con su consecutivo. "
                                        "Queda en espera del XML o para revisión visual."})
    guardar_proveedores(proveedores)
    db.commit()
    return bitacora


def comprobantes_por_clave(db) -> dict:
    return {r["clave"]: r for r in db.consultar("SELECT * FROM comprobantes")}
