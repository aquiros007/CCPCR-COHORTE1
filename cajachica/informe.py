"""Informe consolidado de incumplimientos: todas las unidades de negocio, documento por documento."""
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .db import BaseDatos
from .reglas import REGLAS
from .verificacion import ACCIONES, AVISO_ART_23
from .reporte import AZUL, FMT_FECHA, FMT_MONTO, _pares, _tabla, _titulo

ORDEN = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}


def _rango_gravedad(g: str) -> int:
    return 3 if g.startswith("Muy grave") else 2 if g.startswith("Grave") else 1 if g.startswith("Leve") else 0


def generar_informe(ruta_db: Path, politica: dict, desde: date, hasta: date, destinos: list[Path],
                    incluir_bajas: bool = False) -> list[Path]:
    db = BaseDatos(ruta_db)
    catalogo = politica.get("_catalogo") or {}
    refs = politica.get("referencias_politica") or {}
    sevs = ("ALTA", "MEDIA", "BAJA") if incluir_bajas else ("ALTA", "MEDIA")
    marcas = ",".join("?" * len(sevs))
    rango = (desde.isoformat(), hasta.isoformat())

    liqs = db.consultar("SELECT * FROM liquidaciones WHERE fecha BETWEEN ? AND ?", rango)
    arqs = db.consultar("SELECT * FROM arqueos WHERE fecha BETWEEN ? AND ?", rango)
    hall_fact = db.consultar(f"""
        SELECT h.severidad, h.regla, h.detalle, h.gravedad, h.articulo, f.id AS factura_id, f.caja, f.fila, f.fecha, f.proveedor,
               f.cedula_proveedor, f.consecutivo, f.clave, f.descripcion, f.total_crc, f.estado,
               l.numero AS liquidacion, l.custodio, l.fecha AS fecha_liq, a.nombre AS archivo
        FROM hallazgos h JOIN facturas f ON f.id = h.factura_id
        JOIN liquidaciones l ON l.id = h.liquidacion_id JOIN archivos a ON a.id = h.archivo_id
        WHERE l.fecha BETWEEN ? AND ? AND h.severidad IN ({marcas})""", rango + sevs)
    hall_gen = db.consultar(f"""
        SELECT h.severidad, h.regla, h.detalle, h.gravedad, h.articulo, h.monto, h.caja, h.origen, l.numero AS liquidacion,
               COALESCE(l.custodio, q.custodio) AS custodio, COALESCE(l.fecha, q.fecha) AS fecha, a.nombre AS archivo
        FROM hallazgos h LEFT JOIN liquidaciones l ON l.id = h.liquidacion_id
        LEFT JOIN arqueos q ON q.id = h.arqueo_id JOIN archivos a ON a.id = h.archivo_id
        WHERE h.factura_id IS NULL AND COALESCE(l.fecha, q.fecha) BETWEEN ? AND ? AND h.severidad IN ({marcas})""",
                          rango + sevs)
    verif = db.consultar("""
        SELECT v.caja, v.control, v.articulo, v.estado FROM verificaciones v
        LEFT JOIN liquidaciones l ON l.id = v.liquidacion_id LEFT JOIN arqueos q ON q.id = v.arqueo_id
        WHERE COALESCE(l.fecha, q.fecha) BETWEEN ? AND ?""", rango)
    ult_arqueo = {r["caja"]: r["f"] for r in db.consultar("SELECT caja, MAX(fecha) f FROM arqueos GROUP BY caja")}
    comprobantes = db.consultar("SELECT * FROM comprobantes")
    db.cerrar()
    art = lambda h: h.get("articulo") or refs.get(h["regla"], "")

    unidad = lambda c: (catalogo.get(c) or {}).get("unidad", "")
    responsable = lambda c, alterno="": (catalogo.get(c) or {}).get("responsable") or alterno

    # ---- documentos: una fila por factura, con todos sus incumplimientos
    docs = {}
    for h in hall_fact:
        d = docs.setdefault(h["factura_id"], {**h, "hallazgos": []})
        d["hallazgos"].append(h)
    filas_docs = []
    for d in docs.values():
        hs = sorted(d["hallazgos"], key=lambda x: ORDEN[x["severidad"]])
        filas_docs.append([
            hs[0]["severidad"], d["caja"], unidad(d["caja"]), responsable(d["caja"], d["custodio"]), d["liquidacion"],
            d["archivo"], d["fila"], d["fecha"] and date.fromisoformat(d["fecha"]), d["proveedor"], d["cedula_proveedor"],
            d["consecutivo"], d["clave"], d["descripcion"], d["total_crc"], d["estado"],
            "\n".join(f"• [{x['severidad']}] {REGLAS.get(x['regla'], x['regla'])}: {x['detalle']}" for x in hs),
            max((x["gravedad"] or "" for x in hs), key=_rango_gravedad) or "—",
            "\n".join(sorted({art(x) for x in hs if art(x)})),
            "\n".join(dict.fromkeys(ACCIONES[x["regla"]] for x in hs if x["regla"] in ACCIONES)),
        ])
    filas_docs.sort(key=lambda r: (ORDEN[r[0]], r[1], str(r[4]), r[6] or 0))

    filas_gen = sorted([[
        h["severidad"], h["caja"], unidad(h["caja"]), responsable(h["caja"], h["custodio"]),
        "Arqueo" if h["origen"] == "arqueo" else f"Liquidación {h['liquidacion']}", h["archivo"],
        h["fecha"] and date.fromisoformat(h["fecha"]), REGLAS.get(h["regla"], h["regla"]), h["detalle"], h["monto"],
        h["gravedad"] or "—", art(h), ACCIONES.get(h["regla"], ""),
    ] for h in hall_gen], key=lambda r: (ORDEN[r[0]], r[1]))

    # ---- resumen por unidad
    por_caja = defaultdict(lambda: {"liq": 0, "total": 0, "fact": 0, "rech": 0, "obs": 0, "m_rech": 0, "m_obs": 0,
                                    "alta": 0, "arqueo": None})
    for l in liqs:
        u = por_caja[l["caja"]]
        u["liq"] += 1; u["total"] += l["total_facturas"] or 0; u["fact"] += l["n_facturas"]
        u["rech"] += l["n_rechazadas"]; u["obs"] += l["n_observadas"]
    for d in docs.values():
        u = por_caja[d["caja"]]
        if d["estado"] == "RECHAZADA":
            u["m_rech"] += d["total_crc"] or 0
        elif d["estado"] == "OBSERVADA":
            u["m_obs"] += d["total_crc"] or 0
    for h in hall_fact + hall_gen:
        if h["severidad"] == "ALTA":
            por_caja[h["caja"]]["alta"] += 1
    for a in sorted(arqs, key=lambda a: a["fecha"]):
        por_caja[a["caja"]]["arqueo"] = a
    for c in catalogo:
        por_caja[c]
    filas_unidad = []
    for c, u in por_caja.items():
        if not u["liq"] and not u["arqueo"]:
            semaforo = "SIN ENTREGA"
        elif u["alta"]:
            semaforo = "ROJO"
        elif u["obs"]:
            semaforo = "AMARILLO"
        else:
            semaforo = "VERDE"
        a = u["arqueo"]
        filas_unidad.append([semaforo, c, unidad(c), responsable(c), (catalogo.get(c) or {}).get("fondo"), u["liq"],
                             u["total"], u["fact"], u["rech"], u["m_rech"], u["obs"], u["m_obs"], u["alta"],
                             a and date.fromisoformat(a["fecha"]), a and a["diferencia"]])
    orden_sem = {"ROJO": 0, "AMARILLO": 1, "SIN ENTREGA": 2, "VERDE": 3}
    filas_unidad.sort(key=lambda r: (orden_sem[r[0]], -r[12], r[1]))

    # ---- libro
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, "Informe de incumplimientos — Caja Chica",
            f"Período {desde:%d/%m/%Y} al {hasta:%d/%m/%Y} · Generado {datetime.now():%d/%m/%Y %H:%M}")
    n_cat = len(catalogo) or len(por_caja)
    entregaron = sum(1 for r in filas_unidad if r[0] != "SIN ENTREGA")
    total = sum(r[6] for r in filas_unidad)
    m_rech = sum(r[9] for r in filas_unidad)
    fila = _pares(ws, 4, [
        ("Unidades de negocio", n_cat),
        ("Unidades que entregaron", entregaron),
        ("Unidades sin entrega en el período", sum(1 for r in filas_unidad if r[0] == "SIN ENTREGA")),
        ("Unidades en ROJO (incumplimientos graves)", sum(1 for r in filas_unidad if r[0] == "ROJO")),
        ("Liquidaciones revisadas", len(liqs)),
        ("Arqueos revisados", len(arqs)),
        ("Gasto presentado", total, FMT_MONTO),
        ("Monto que NO debe reintegrarse (rechazado)", m_rech, FMT_MONTO),
        ("Monto pendiente de justificar (observado)", sum(r[11] for r in filas_unidad), FMT_MONTO),
        ("Documentos con incumplimientos", len(filas_docs)),
        ("   de severidad ALTA", sum(1 for r in filas_docs if r[0] == "ALTA")),
        ("Incumplimientos de liquidación o arqueo", len(filas_gen)),
    ])
    ws.cell(row=fila + 1, column=1, value="Incumplimientos más frecuentes").font = Font(bold=True, color=AZUL)
    conteo = defaultdict(lambda: [0, 0.0])
    for h in hall_fact:
        conteo[h["regla"]][0] += 1; conteo[h["regla"]][1] += h["total_crc"] or 0
    for h in hall_gen:
        conteo[h["regla"]][0] += 1; conteo[h["regla"]][1] += abs(h["monto"] or 0)
    top = sorted(conteo.items(), key=lambda x: -x[1][0])[:15]
    for i, (regla, (n, m)) in enumerate(top, fila + 2):
        ws.cell(row=i, column=1, value=REGLAS.get(regla, regla))
        ws.cell(row=i, column=2, value=f"{n} casos · ₡{m:,.0f}")
    ws.cell(row=fila + len(top) + 3, column=1, value=AVISO_ART_23).font = Font(italic=True, color="666666")

    # ---- lista de verificación consolidada: cuántos documentos cumplen cada control
    ws = wb.create_sheet("Lista de verificación")
    por_control = defaultdict(lambda: defaultdict(int))
    for v in verif:
        por_control[(v["control"], v["articulo"])][v["estado"]] += 1
    estados = ["CUMPLE", "NO CUMPLE", "NO SE PUEDE VERIFICAR", "NO APLICA"]
    _tabla(ws, 1, ["Control", "Artículo"] + estados,
           [[c, a] + [por_control[(c, a)][e] for e in estados] for (c, a) in por_control],
           [60, 26, 11, 12, 14, 11])

    ws = wb.create_sheet("Por unidad")
    _tabla(ws, 1, ["Semáforo", "Código", "Unidad de negocio", "Responsable", "Fondo", "Liquidaciones", "Gasto presentado",
                   "Facturas", "Rechazadas", "Monto rechazado", "Observadas", "Monto observado", "Hallazgos ALTA",
                   "Último arqueo", "Diferencia arqueo"], filas_unidad,
           [13, 11, 28, 24, 13, 8, 15, 8, 9, 15, 9, 15, 9, 12, 14], montos=(5, 7, 10, 12, 15), fechas=(14,))
    colores = {"ROJO": "F8D7DA", "AMARILLO": "FFF3CD", "VERDE": "D4EDDA", "SIN ENTREGA": "E2E3E5"}
    for r in range(2, len(filas_unidad) + 2):
        c = ws.cell(row=r, column=1)
        c.fill = PatternFill("solid", fgColor=colores[c.value]); c.font = Font(bold=True)

    ws = wb.create_sheet("Documentos incumplidos")
    _tabla(ws, 1, ["Severidad", "Código", "Unidad", "Responsable", "Liquidación", "Archivo", "Fila", "Fecha factura",
                   "Proveedor", "Cédula proveedor", "N° factura", "Clave numérica", "Descripción", "Monto", "Estado",
                   "Incumplimientos", "Gravedad sugerida", "Artículo", "Acción requerida"], filas_docs,
           [10, 11, 22, 20, 12, 30, 6, 12, 26, 14, 22, 20, 30, 12, 12, 70, 20, 26, 40],
           col_estado=1, montos=(14,), fechas=(8,))

    ws = wb.create_sheet("Liquidaciones y arqueos")
    _tabla(ws, 1, ["Severidad", "Código", "Unidad", "Responsable", "Documento", "Archivo", "Fecha", "Incumplimiento",
                   "Detalle", "Monto", "Gravedad sugerida", "Artículo", "Acción requerida"], filas_gen,
           [10, 11, 22, 20, 18, 30, 12, 34, 70, 13, 20, 26, 40], col_estado=1, montos=(10,), fechas=(7,))

    sin = [[r[1], r[2], r[3], (catalogo.get(r[1]) or {}).get("correo", ""), r[4]] for r in filas_unidad if r[0] == "SIN ENTREGA"]
    ws = wb.create_sheet("Sin entrega")
    _tabla(ws, 1, ["Código", "Unidad de negocio", "Responsable", "Correo", "Fondo"], sin, [11, 30, 26, 30, 14], montos=(5,))

    # ---- apertura y catálogo (arts. 4, 5, 13 b)
    from .config import problemas_catalogo
    ws = wb.create_sheet("Apertura y catálogo")
    _tabla(ws, 1, ["FUC", "Problema", "Artículo"],
           [[p["codigo"], p["problema"], p["articulo"]] for p in problemas_catalogo(politica)], [24, 90, 16])

    # ---- arqueos periódicos (arts. 14 y 17)
    maximo = int((politica.get("plazos") or {}).get("dias_maximos_entre_arqueos", 30))
    sin_arqueo = []
    for c in sorted(set(catalogo) | set(ult_arqueo)):
        ultimo = ult_arqueo.get(c)
        dias = (hasta - date.fromisoformat(ultimo[:10])).days if ultimo else None
        if dias is None or dias > maximo:
            sin_arqueo.append([c, unidad(c), responsable(c), ultimo and date.fromisoformat(ultimo[:10]),
                               "Nunca" if dias is None else f"{dias} días"])
    ws = wb.create_sheet("Sin arqueo reciente")
    _tabla(ws, 1, ["Código", "Unidad de negocio", "Responsable", "Último arqueo", f"Sin arqueo (máx. {maximo} días)"],
           sin_arqueo, [11, 30, 26, 14, 22], fechas=(4,))

    # ---- comprobantes electrónicos de proveedores
    ws = wb.create_sheet("Comprobantes proveedores")
    import json as _json
    filas_comp = [[c["estado"], c["emisor_nombre"], c["emisor_cedula"], c["consecutivo"], c["fecha"], c["total"],
                   c["estado_hacienda"], "\n".join(f"{x['control']}: {x['detalle']}" for x in _json.loads(c["controles"])
                                                  if x["estado"] != "CUMPLE"), c["clave"]]
                  for c in sorted(comprobantes, key=lambda c: c["estado"])]
    _tabla(ws, 1, ["Estado", "Proveedor", "Cédula", "N° factura", "Fecha", "Total", "Hacienda", "Controles pendientes",
                   "Clave"], filas_comp, [22, 30, 14, 22, 11, 13, 14, 70, 54], montos=(6,))

    escritos = []
    for d in destinos:
        d.parent.mkdir(parents=True, exist_ok=True)
        wb.save(d)
        escritos.append(d)
    return escritos
